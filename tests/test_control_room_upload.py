"""Control Room upload, match, review, and manifest tests; no network."""
import hashlib
import io
import sqlite3
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from control_room.app import create_app
from control_room.fake import FakeEnrichmentAdapter
from control_room.matching import CanonicalMatcher, normalize_name, normalize_tax
from control_room.repository import JobRepository
from control_room.uploads import UploadError, parse_csv, parse_xlsx, suggest_mapping
from control_room.worker import Worker


def canonical_database(path):
    conn = sqlite3.connect(path)
    conn.executescript('''CREATE TABLE companies_lite(
        id INTEGER PRIMARY KEY, company_name TEXT, registration_number TEXT,
        tax_number TEXT, address TEXT, municipality TEXT);
        INSERT INTO companies_lite VALUES
        (0,'ELBI, d.o.o.','5483883','32670338','Cesta 1','4000 Kranj'),
        (1,'ALFA d.o.o.','1111111','12345678','Glavna cesta 1','Kranj'),
        (2,'BETA d.o.o.','2222222','87654321','Trg 2','Ljubljana'),
        (3,'DVOJNIK d.o.o.','3333333','11112222','Pot 3','Celje'),
        (4,'DVOJNIK d.o.o.','4444444','33334444','Pot 4','Maribor'),
        (5,'BAKRA d.o.o. Kranj','5557755','58626450','Cesta 5','4000 Kranj');''')
    conn.close()


class ParserAndMatcherTests(unittest.TestCase):
    @staticmethod
    def zip_file(path, entries, compression=zipfile.ZIP_DEFLATED):
        with zipfile.ZipFile(path, 'w', compression=compression) as archive:
            for name, content in entries:
                archive.writestr(name, content)

    def test_csv_utf8_bom_semicolon_and_tab(self):
        headers, rows, sheet = parse_csv('\ufeffNaziv;Davčna številka\nALFA;SI12345678\n'.encode(), 10)
        self.assertEqual((headers, rows, sheet),
            (['Naziv','Davčna številka'],[['ALFA','SI12345678']],None))
        self.assertEqual(parse_csv(b'Naziv\tObcina\nALFA\tKranj\n',10)[1][0],['ALFA','Kranj'])

    def test_csv_rejects_row_limit_and_malformed_width(self):
        with self.assertRaisesRegex(UploadError,'row limit'):
            parse_csv(b'Name,Tax\nOne,1\nTwo,2\n',1)
        with self.assertRaises(UploadError):
            parse_csv(b'Name,Tax\nOne,1,extra\n',10)

    def test_xlsx_first_visible_sheet_data_only(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)/'companies.xlsx'; workbook = Workbook()
            hidden = workbook.active; hidden.title='Hidden'; hidden.sheet_state='hidden'
            sheet = workbook.create_sheet('Companies'); sheet.append(['Naziv','Davčna'])
            sheet.append(['ALFA','=1+1']); workbook.save(path)
            headers, rows, title = parse_xlsx(path,10)
            self.assertEqual((headers,title),(['Naziv','Davčna'],'Companies'))
            self.assertEqual(rows[0][0],'ALFA')
            self.assertEqual(rows[0][1],'')  # formulas are never evaluated

    def test_xlsx_pads_trailing_cells_and_preserves_interstitial_rows(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)/'registrations.xlsx'; workbook = Workbook()
            sheet = workbook.active; sheet.title = 'Registrations'
            sheet.append(['Ime', 'Email', 'Podjetje', 'Source'])
            sheet.append(['Ana', 'ana@example.si'])
            sheet.append([None, None, None, None])
            sheet.append(['Bine', None, 'BETA'])
            sheet.append([None, None, None, None])  # unused trailing row
            workbook.save(path)
            headers, rows, title = parse_xlsx(path, 10)
            self.assertEqual(title, 'Registrations')
            self.assertEqual(headers, ['Ime', 'Email', 'Podjetje', 'Source'])
            self.assertEqual(rows, [
                ['Ana', 'ana@example.si', '', ''],
                ['', '', '', ''],
                ['Bine', '', 'BETA', ''],
            ])

    def test_xlsx_rejects_cells_beyond_header_width(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)/'malformed.xlsx'; workbook = Workbook()
            sheet = workbook.active
            sheet.append(['Ime', 'Podjetje'])
            sheet.append(['Ana', 'ALFA', 'unexpected'])
            workbook.save(path)
            with self.assertRaisesRegex(UploadError, 'exceeds the header width'):
                parse_xlsx(path, 10)

    def test_xlsx_rejects_excessive_zip_entry_count_before_parsing(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)/'many-entries.xlsx'
            self.zip_file(path, [('entry-%d.xml' % index, b'x') for index in range(3)])
            with patch('control_room.uploads.XLSX_MAX_ZIP_ENTRIES', 2):
                with self.assertRaisesRegex(UploadError, 'too complex'):
                    parse_xlsx(path, 10)

    def test_xlsx_rejects_excessive_total_declared_size_before_parsing(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)/'large-total.xlsx'
            self.zip_file(path, [('one.xml', b'123456'), ('two.xml', b'123456')])
            with patch('control_room.uploads.XLSX_MAX_UNCOMPRESSED_BYTES', 10), \
                    patch('control_room.uploads.XLSX_MAX_ENTRY_BYTES', 10):
                with self.assertRaisesRegex(UploadError, 'too complex'):
                    parse_xlsx(path, 10)

    def test_xlsx_rejects_excessive_single_entry_before_parsing(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)/'large-entry.xlsx'
            self.zip_file(path, [('large.xml', b'12345')])
            with patch('control_room.uploads.XLSX_MAX_ENTRY_BYTES', 4):
                with self.assertRaisesRegex(UploadError, 'too complex'):
                    parse_xlsx(path, 10)

    def test_xlsx_rejects_suspicious_compression_ratio_before_parsing(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)/'compressed.xlsx'
            self.zip_file(path, [('repeated.xml', b'A' * 4096)])
            with patch('control_room.uploads.XLSX_COMPRESSION_RATIO_MIN_BYTES', 1), \
                    patch('control_room.uploads.XLSX_MAX_COMPRESSION_RATIO', 2):
                with self.assertRaisesRegex(UploadError, 'too complex'):
                    parse_xlsx(path, 10)

    def test_xlsx_rejects_malformed_zip_safely_and_csv_skips_zip_validation(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)/'malformed.xlsx'; path.write_bytes(b'not a zip')
            with self.assertRaisesRegex(UploadError, 'Invalid XLSX workbook'):
                parse_xlsx(path, 10)
        with patch('control_room.uploads.validate_xlsx_archive',
                   side_effect=AssertionError('CSV must not inspect ZIP metadata')):
            self.assertEqual(parse_csv(b'Name,Tax\nALFA,1\n', 10)[1], [['ALFA', '1']])

    def test_slovenian_mapping_identifiers_only_and_ambiguous_headers(self):
        mapping = suggest_mapping(['Naziv podjetja','Davčna številka','Matična','Naslov','Občina'])
        self.assertEqual(mapping,{'company_name':0,'tax_number':1,'registration_number':2,
                                  'address':3,'municipality':4})
        identifiers = suggest_mapping(['Davcna stevilka','Maticna stevilka'])
        self.assertEqual((identifiers['tax_number'],identifiers['registration_number']),(0,1))
        self.assertIsNone(suggest_mapping(['Naziv','Podjetje'])['company_name'])

    def test_matching_order_and_conservative_fuzzy(self):
        companies = [
            {'id':1,'company_name':'ALFA d.o.o.','tax_number':'12345678','registration_number':'1111111','address':'Glavna 1','municipality':'Kranj'},
            {'id':2,'company_name':'DVOJNIK d.o.o.','tax_number':'22222222','registration_number':'2222222','address':'A 1','municipality':'Celje'},
            {'id':3,'company_name':'DVOJNIK d.o.o.','tax_number':'33333333','registration_number':'3333333','address':'B 2','municipality':'Maribor'}]
        matcher = CanonicalMatcher(companies)
        def row(name='',tax='',registration='',address='',municipality=''):
            return {'normalized_name':normalize_name(name),'normalized_tax_number':normalize_tax(tax),
                'normalized_registration_number':''.join(filter(str.isdigit,registration)),
                'normalized_address':address.lower(),'normalized_municipality':municipality.lower()}
        self.assertEqual(matcher.match(row(tax='SI 12345678'))[1][0]['match_method'],'TAX_EXACT')
        self.assertEqual(matcher.match(row(registration='1111111'))[1][0]['match_method'],'REGISTRATION_EXACT')
        self.assertEqual(matcher.match(row('DVOJNIK','', '',municipality='Celje'))[1][0]['match_method'],'NAME_LOCATION')
        self.assertEqual(matcher.match(row('ALFA'))[0],'MATCHED')
        self.assertEqual(matcher.match(row('DVOJNIK'))[0],'AMBIGUOUS')
        status,candidates = matcher.match(row('ALFAA'))
        self.assertEqual((status,candidates[0]['match_method']),('AMBIGUOUS','FUZZY_CANDIDATE'))
        self.assertEqual(matcher.match(row('NONEXISTENT XYZZY'))[0],'NOT_FOUND')


class UploadWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.root = Path(self.temp.name)
        self.canonical = self.root/'canonical.sqlite3'; canonical_database(self.canonical)
        self.control = self.root/'control.sqlite3'; self.storage = self.root/'storage'
        self.config = {'TESTING':True,'CONTROL_DB':self.control,'CANONICAL_DB':self.canonical,
            'CONTROL_STORAGE_ROOT':self.storage,'DISCOVERY_V2_PATH':self.root/'missing-v2',
            'UPLOAD_MAX_BYTES':1024*1024,'UPLOAD_MAX_ROWS':20}
        self.app = create_app(self.config); self.client = self.app.test_client()
        self.repository = self.app.extensions['control_repository']

    def tearDown(self): self.temp.cleanup()

    def upload(self, content=None, filename='companies.csv'):
        content = content or b'Naziv podjetja;Davcna stevilka;Obcina\nALFA d.o.o.;SI12345678;Kranj\n'
        return self.client.post('/uploads',data={'company_file':(io.BytesIO(content),filename)},
                                content_type='multipart/form-data')

    def mapped_upload(self, content=None):
        response = self.upload(content); self.assertEqual(response.status_code,303)
        upload_id = response.headers['Location'].split('/')[-2]
        response = self.client.post(f'/uploads/{upload_id}/mapping',data={
            'company_name':'0','tax_number':'1','registration_number':'','address':'','municipality':'2'})
        self.assertEqual(response.status_code,303)
        return upload_id

    def test_upload_generated_path_hash_and_canonical_read_only(self):
        before = hashlib.sha256(self.canonical.read_bytes()).hexdigest()
        response = self.upload(filename='../../evil.csv'); self.assertEqual(response.status_code,303)
        upload_id = response.headers['Location'].split('/')[-2]
        upload = self.repository.get_upload(upload_id)
        self.assertEqual(upload['original_filename'],'evil.csv')
        self.assertEqual(Path(upload['relative_path']).parent,Path('uploads'))
        self.assertNotIn('evil',Path(upload['relative_path']).name)
        stored = self.storage/upload['relative_path']; self.assertTrue(stored.is_file())
        self.assertEqual(stored.stat().st_mode & 0o777,0o600)
        self.assertEqual(upload['sha256'],hashlib.sha256(stored.read_bytes()).hexdigest())
        self.assertEqual(before,hashlib.sha256(self.canonical.read_bytes()).hexdigest())

    def test_invalid_extension_oversize_row_limit_and_malformed(self):
        self.assertEqual(self.upload(b'x','bad.xls').status_code,400)
        app = create_app({**self.config,'CONTROL_DB':self.root/'small.sqlite3',
                          'CONTROL_STORAGE_ROOT':self.root/'small','UPLOAD_MAX_BYTES':8})
        self.assertEqual(app.test_client().post('/uploads',data={
            'company_file':(io.BytesIO(b'Name\n123456789\n'),'x.csv')},content_type='multipart/form-data').status_code,400)
        rows = b'Name\n'+b'X\n'*21
        self.assertEqual(self.upload(rows).status_code,400)
        self.assertEqual(self.upload(b'Name,Tax\nOne,1,extra\n').status_code,400)

    def test_xlsx_upload(self):
        workbook=Workbook(); sheet=workbook.active; sheet.append(['Naziv','Davcna']); sheet.append(['ALFA','12345678'])
        stream=io.BytesIO(); workbook.save(stream)
        response=self.upload(stream.getvalue(),'companies.xlsx'); self.assertEqual(response.status_code,303)
        upload=self.repository.get_upload(response.headers['Location'].split('/')[-2])
        self.assertEqual((upload['format'],upload['worksheet_name']),('XLSX','Sheet'))

    def test_manual_remapping_and_identifier_only_file(self):
        response=self.upload(b'Wrong;Davcna stevilka\nIgnore;SI12345678\n')
        upload_id=response.headers['Location'].split('/')[-2]
        response=self.client.post(f'/uploads/{upload_id}/mapping',data={
            'company_name':'','tax_number':'1','registration_number':'','address':'','municipality':''})
        self.assertEqual(response.status_code,303)
        row=self.repository.upload_rows(upload_id)[0]
        self.assertEqual((row['match_status'],row['selected_company_id'],row['match_method']),
                         ('MATCHED',1,'TAX_EXACT'))

    def test_review_resolution_exclusion_and_recreation(self):
        upload_id=self.mapped_upload(b'Naziv;Davcna;Obcina\nDVOJNIK;;\nNIMA PODJETJA;;\n')
        rows=self.repository.upload_rows(upload_id)
        self.assertEqual([row['match_status'] for row in rows],['AMBIGUOUS','NOT_FOUND'])
        candidates=self.repository.candidates(upload_id)
        response=self.client.post(f'/uploads/{upload_id}/rows/2',data={
            'action':'select','company_id':str(candidates[2][0]['company_id'])})
        self.assertEqual(response.status_code,303)
        self.client.post(f'/uploads/{upload_id}/rows/3',data={'action':'exclude'})
        recreated=create_app(self.config).extensions['control_repository']
        persisted=recreated.upload_rows(upload_id)
        self.assertEqual([row['match_status'] for row in persisted],['MATCHED','EXCLUDED'])

    def test_manual_candidate_post_persists_bakra_and_valid_company_id_zero(self):
        content = ('Naziv;Davcna;Obcina\n'
                   'ELBI, d.o.o.;SI32670338;4000 Kranj\n'
                   'BAKRA d.o.o. Kranj X;;4000 Kranj\n'
                   'fake company;;\n').encode()
        upload_id = self.mapped_upload(content)
        rows = self.repository.upload_rows(upload_id)
        self.assertEqual((rows[0]['selected_company_id'], rows[0]['match_status']), (0, 'MATCHED'))
        candidates = self.repository.candidates(upload_id)
        bakra_id = candidates[3][0]['company_id']
        response = self.client.post(
            f'/uploads/{upload_id}/rows/3/select/{bakra_id}')
        self.assertEqual(response.status_code, 303)
        self.assertEqual(self.client.post(f'/uploads/{upload_id}/rows/4',
            data={'action':'exclude'}).status_code, 303)
        recreated = create_app(self.config).extensions['control_repository']
        persisted = recreated.upload_rows(upload_id)
        self.assertEqual((persisted[1]['match_status'], persisted[1]['selected_company_id'],
                          persisted[1]['selected'], persisted[1]['match_method']),
                         ('MATCHED', bakra_id, 1, 'MANUAL_FUZZY_CANDIDATE'))
        self.assertEqual(persisted[2]['match_status'], 'EXCLUDED')
        self.assertEqual(recreated.upload_summary(upload_id)['selected_unique'], 2)
        selected = recreated.candidates(upload_id)[3]
        self.assertEqual([item['selected'] for item in selected], [1])

        zero_upload = self.mapped_upload(b'Naziv;Davcna;Obcina\nELBII;;4000 Kranj\n')
        zero_candidate = self.repository.candidates(zero_upload)[2][0]
        self.assertEqual(zero_candidate['company_id'], 0)
        response = self.client.post(f'/uploads/{zero_upload}/rows/2/select/0')
        self.assertEqual(response.status_code, 303)
        zero_row = self.repository.upload_rows(zero_upload)[0]
        self.assertEqual((zero_row['match_status'],zero_row['selected_company_id'],
                          zero_row['selected'],zero_row['match_method']),
                         ('MATCHED',0,1,'MANUAL_FUZZY_CANDIDATE'))

    def test_duplicate_rows_manifest_is_immutable_and_fake_count_is_unique(self):
        upload_id=self.mapped_upload(b'Naziv;Davcna;Obcina\nALFA;12345678;Kranj\nALFA;12345678;Kranj\nBETA;87654321;Ljubljana\n')
        response=self.client.post(f'/uploads/{upload_id}/confirm',data={'display_name':'Upload job'})
        self.assertEqual(response.status_code,303)
        job=self.repository.list_jobs()[0]; items=self.repository.job_items(job['job_id'])
        self.assertEqual((job['input_kind'],job['selected_company_count'],len(items)),('UPLOAD',2,3))
        self.assertEqual(sum(item['selected'] for item in items),2)
        self.assertEqual(len(self.repository.upload_rows(upload_id)),3)
        with self.assertRaisesRegex(ValueError,'closed'):
            self.repository.resolve_upload_row(upload_id,2)
        self.assertEqual([item['company_id'] for item in self.repository.job_items(job['job_id'])],[1,1,2])
        Worker(self.repository,FakeEnrichmentAdapter(delay=0,sleeper=lambda _:None),worker_id='upload-worker').run_once()
        self.assertEqual((self.repository.get_job(job['job_id'])['status'],
                          self.repository.get_job(job['job_id'])['processed_company_count']),('COMPLETED',2))

    def test_pages_escape_input_and_workflow_survives_reload(self):
        upload_id=self.mapped_upload(b'Naziv;Davcna;Obcina\n<script>alert(1)</script>;;\n')
        html=self.client.get(f'/uploads/{upload_id}/review').get_data(as_text=True)
        self.assertIn('&lt;script&gt;',html); self.assertNotIn('<script>alert(1)</script>',html)
        second=create_app(self.config).test_client()
        self.assertEqual(second.get(f'/uploads/{upload_id}/review').status_code,200)
        self.assertEqual(second.get(f'/uploads/{upload_id}/confirm').status_code,200)

    def test_v1_database_migrates_without_losing_job_or_event(self):
        old=self.root/'old.sqlite3'; schema=Path('control_room/schema.sql').read_text().replace(
            'BETWEEN 1 AND 5000','BETWEEN 1 AND 1000').replace('PRAGMA user_version=2','PRAGMA user_version=1')
        conn=sqlite3.connect(old); conn.executescript(schema)
        conn.execute("INSERT INTO control_jobs(job_id,display_name,input_kind,module,status,created_at,selected_company_count) VALUES ('old','Old','FAKE','DISCOVERY_CONTACTS','QUEUED','now',1)")
        conn.execute("INSERT INTO job_events(job_id,sequence,created_at,level,event_code,message) VALUES ('old',1,'now','INFO','OLD','old event')")
        conn.commit(); conn.close()
        repository=JobRepository(old); repository.initialize()
        self.assertEqual(repository.get_job('old')['display_name'],'Old')
        self.assertEqual(repository.events('old')[0]['event_code'],'OLD')
        self.assertEqual(repository.get_job('old')['execution_adapter'],'FAKE')


if __name__ == '__main__': unittest.main()
