"""Customer-facing request-aware MERLIN XLSX export tests."""
import io
import sqlite3
from unittest.mock import patch

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from werkzeug.datastructures import FileStorage

from control_room.import_enrich_adapter import ImportEnrichAdapter
from control_room.import_export import (canonical_merlin_website,
                                        export_import_xlsx)
from control_room.repository import JobRepository
from control_room.worker import Worker
from merlin.service import MerlinImportEnrichService


def _canonical(path):
    conn = sqlite3.connect(path)
    conn.executescript('''
      CREATE TABLE companies_lite(
        id INTEGER PRIMARY KEY,company_name TEXT,registration_number TEXT,
        tax_number TEXT,address TEXT,municipality TEXT,revenue_2025 REAL,
        profit_2025 REAL,employees_2025 REAL,assets_2025 REAL,capital_2025 REAL);
      INSERT INTO companies_lite VALUES
        (1,'ALFA d.o.o.','1001','11111111','Alfa 1','Kranj',100,12,3,250,NULL),
        (2,'DVOJNIK d.o.o.','1002','22222222','Prva 2','Celje',200,22,4,350,20),
        (3,'DVOJNIK d.o.o.','1003','33333333','Druga 3','Maribor',300,32,5,450,30);
      CREATE TABLE website_discovery(
        id INTEGER PRIMARY KEY,company_id INTEGER,website TEXT,status TEXT,
        verified_scope TEXT,relationship TEXT);
      INSERT INTO website_discovery VALUES
        (1,1,'https://alfa.si/','VERIFIED','https://alfa.si/','LEGAL_ENTITY');
      CREATE TABLE email_discovery(
        id INTEGER PRIMARY KEY,company_id INTEGER,email TEXT,website TEXT);
      INSERT INTO email_discovery VALUES
        (1,1,'legacy@alfa.si','https://alfa.si/');
    ''')
    conn.commit(); conn.close()


def _accepted_contacts(path, website='https://alfa.si/'):
    conn = sqlite3.connect(path)
    conn.executescript('''
      CREATE TABLE discovery_runs(
        run_id TEXT PRIMARY KEY,status TEXT,engine_version TEXT,rule_version TEXT);
      CREATE TABLE discovery_run_companies(
        run_id TEXT,company_id INTEGER,status TEXT,selected_attempt_id TEXT);
      CREATE TABLE discovery_company_results(
        result_id TEXT,run_id TEXT,attempt_id TEXT,company_id INTEGER,
        website_status TEXT,official_website TEXT,website_observation_id TEXT,
        website_evidence_ids_json TEXT,default_email_contact_id TEXT,
        default_phone_contact_id TEXT,contact_outcome TEXT,completed_at TEXT);
      CREATE TABLE discovery_contacts(
        contact_id TEXT,run_id TEXT,attempt_id TEXT,company_id INTEGER,
        contact_type TEXT,normalized_value TEXT,primary_observation_id TEXT,
        supporting_observation_ids_json TEXT,attribution_status TEXT,
        rule_version TEXT,roles_json TEXT);
      INSERT INTO discovery_runs VALUES ('accepted','COMPLETED','engine','rules');
      INSERT INTO discovery_run_companies VALUES ('accepted',1,'COMPLETED','attempt');
      INSERT INTO discovery_company_results VALUES
        ('result','accepted','attempt',1,'HIGH',NULL,'website',
         '["website"]','email','phone','FOUND','2026-10-01T00:00:00+00:00');
      INSERT INTO discovery_contacts VALUES
        ('email','accepted','attempt',1,'EMAIL','selected@alfa.si','email-observation',
         '["email-observation"]','ATTRIBUTED','rules','["GENERAL"]'),
        ('phone','accepted','attempt',1,'PHONE','+38611234567','phone-observation',
         '["phone-observation"]','ATTRIBUTED','rules','["MAIN"]');
    ''')
    conn.execute('UPDATE discovery_company_results SET official_website=?', (website,))
    conn.commit(); conn.close()


def _source_rows():
    return [
        ['ALFA d.o.o.', '11111111', '=1+1'],
        ['ALFA d.o.o.', '11111111', '=1+1'],
        ['DVOJNIK d.o.o.', '', 'ambiguous'],
        ['NEZNANO d.o.o.', '', 'not found'],
    ]


def _xlsx_upload(collision=False):
    stream = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Source data'
    headers = ['Company', 'Tax', 'Source value']
    if collision:
        headers.append('Website')
    sheet.append(headers)
    for values in _source_rows():
        sheet.append(values + (['customer website'] if collision else []))
    sheet['A1'].font = Font(bold=True, color='FFFFFF')
    sheet['A1'].fill = PatternFill('solid', fgColor='204060')
    sheet.row_dimensions[2].height = 27
    sheet.column_dimensions['A'].width = 24
    hidden = workbook.create_sheet('Hidden source')
    hidden.sheet_state = 'hidden'; hidden['A1'] = 'preserved'
    merged = workbook.create_sheet('Merged source')
    merged.merge_cells('A1:C1'); merged['A1'] = 'preserved merge'
    workbook.save(stream); workbook.close(); stream.seek(0)
    return FileStorage(stream=stream, filename='source.xlsx',
                       content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def _csv_upload():
    data = ('Company,Tax,Source value\n'
            'ALFA d.o.o.,11111111,"=1+1"\n'
            'ALFA d.o.o.,11111111,"=1+1"\n'
            'DVOJNIK d.o.o.,,ambiguous\n'
            'NEZNANO d.o.o.,,not found\n').encode()
    return FileStorage(stream=io.BytesIO(data), filename='source.csv',
                       content_type='text/csv')


def _run(tmp_path, requested_outputs, *, source_format='XLSX', collision=False,
         website='https://alfa.si/'):
    root = tmp_path/'storage'
    canonical = tmp_path/'canonical.sqlite3'; _canonical(canonical)
    contacts = tmp_path/'contacts.sqlite3'; _accepted_contacts(contacts, website)
    repository = JobRepository(tmp_path/'control.sqlite3'); repository.initialize()
    service = MerlinImportEnrichService(repository, root)
    workspace = service.create_workspace('Workspace')['workspace_id']
    upload = service.ingest_upload(workspace,
        _xlsx_upload(collision) if source_format == 'XLSX' else _csv_upload())
    job = service.queue_import(workspace, upload.upload_id, 'Customer export',
        {'company_name': 0, 'tax_number': 1}, requested_outputs)
    adapter = ImportEnrichAdapter(repository, canonical, (contacts,), storage_root=root)
    with patch('requests.sessions.Session.request',
               side_effect=AssertionError('network/provider call forbidden')):
        worker = Worker(repository, adapter, worker_id='merlin-export',
                        adapters={'IMPORT_ENRICH': adapter})
        assert worker.run_once()
    completed = service.job(workspace, job.job_id)
    assert completed.status == 'COMPLETED'
    stored = repository.artifacts(job.job_id)
    assert len(stored) == 1 and stored[0]['artifact_type'] == 'UPLOAD_RECONCILIATION'
    return root, repository, service, workspace, upload, completed, stored[0]


@pytest.mark.parametrize(('outputs', 'expected'), [
    (('WEBSITE',), ['Merlin match', 'Merlin note', 'Website']),
    (('EMAIL', 'PHONE'), ['Merlin match', 'Merlin note', 'Email', 'Phone']),
    (('FINANCIALS',), ['Merlin match', 'Merlin note', 'Revenue 2025',
        'Profit 2025', 'Employees 2025', 'Assets 2025', 'Capital 2025']),
    (('WEBSITE', 'EMAIL', 'PHONE', 'FINANCIALS'), ['Merlin match', 'Merlin note',
        'Website', 'Email', 'Phone', 'Revenue 2025', 'Profit 2025',
        'Employees 2025', 'Assets 2025', 'Capital 2025']),
])
def test_merlin_xlsx_exports_only_requested_customer_columns(tmp_path, outputs, expected):
    root, _, _, _, _, _, artifact = _run(tmp_path, outputs)
    workbook = load_workbook(root/artifact['relative_path'], data_only=False)
    try:
        sheet = workbook['Source data']
        headers = [cell.value for cell in sheet[1]]
        assert headers[:3] == ['Company', 'Tax', 'Source value']
        assert headers[3:] == expected
        lookup = {value: index for index, value in enumerate(headers, 1)}
        if 'Email' in expected:
            assert sheet.cell(2, lookup['Email']).value == 'selected@alfa.si'
        if 'Phone' in expected:
            assert sheet.cell(2, lookup['Phone']).value == '+38611234567'
        if 'Website' not in expected:
            assert 'Website' not in headers
        if outputs == ('FINANCIALS',):
            assert [sheet.cell(2, lookup[name]).value for name in expected[2:]] == [
                100, 12, 3, 250, None]
            assert all(isinstance(sheet.cell(2, lookup[name]).value, (int, float))
                       for name in expected[2:-1])
    finally:
        workbook.close()


@pytest.mark.parametrize(('stored', 'customer'), [
    ('https://example.si/kontakt/', 'https://example.si/'),
    ('https://example.si/about?q=company#team', 'https://example.si/'),
    ('https://www.example.si/company', 'https://www.example.si/'),
    ('http://example.si/company', 'http://example.si/'),
    ('https://example.si:8443/company', 'https://example.si:8443/'),
    ('not a valid website', 'not a valid website'),
    ('https://example.si:invalid/company', 'https://example.si:invalid/company'),
    ('ftp://example.si/company', 'ftp://example.si/company'),
])
def test_merlin_customer_website_projection_is_conservative(stored, customer):
    assert canonical_merlin_website(stored) == customer


def test_merlin_export_projects_website_root_without_mutating_stored_evidence(tmp_path):
    evidence_url = 'https://www.example.si/kontakt/?from=directory#office'
    root, repository, _, _, _, job, artifact = _run(
        tmp_path, ('WEBSITE',), website=evidence_url)
    workbook = load_workbook(root/artifact['relative_path'], data_only=False)
    try:
        sheet = workbook['Source data']
        headers = [cell.value for cell in sheet[1]]
        website_column = headers.index('Website') + 1
        assert sheet.cell(2, website_column).value == 'https://www.example.si/'
    finally:
        workbook.close()

    stored_item = repository.job_items(job.job_id)[0]
    assert stored_item['enrichment']['official_website'] == evidence_url
    connection = sqlite3.connect(tmp_path/'contacts.sqlite3')
    try:
        assert connection.execute(
            'SELECT official_website FROM discovery_company_results').fetchone()[0] == evidence_url
    finally:
        connection.close()


def test_email_only_merlin_export_does_not_leak_stored_website(tmp_path):
    evidence_url = 'https://example.si/contact/'
    root, _, _, _, _, _, artifact = _run(
        tmp_path, ('EMAIL',), website=evidence_url)
    workbook = load_workbook(root/artifact['relative_path'], data_only=False)
    try:
        sheet = workbook['Source data']
        headers = [cell.value for cell in sheet[1]]
        assert 'Website' not in headers
        assert evidence_url not in [cell.value for row in sheet.iter_rows() for cell in row]
    finally:
        workbook.close()


def test_merlin_collision_and_lossless_xlsx_guarantees(tmp_path):
    root, repository, _, _, upload, _, artifact = _run(
        tmp_path, ('WEBSITE',), collision=True)
    stored_upload = repository.get_upload(upload.upload_id)
    source = root/stored_upload['relative_path']
    original = load_workbook(source, data_only=False)
    exported = load_workbook(root/artifact['relative_path'], data_only=False)
    try:
        assert exported.sheetnames == [
            'Source data', 'Hidden source', 'Merged source', 'Merlin review']
        old, new = original['Source data'], exported['Source data']
        assert new['C2'].value == old['C2'].value == '=1+1'
        assert new['A1'].font.bold == old['A1'].font.bold is True
        assert new['A1'].fill.fgColor.rgb == old['A1'].fill.fgColor.rgb
        assert new.row_dimensions[2].height == 27
        assert new.column_dimensions['A'].width == 24
        assert exported['Hidden source'].sheet_state == 'hidden'
        assert list(exported['Merged source'].merged_cells.ranges)[0].coord == 'A1:C1'
        headers = [cell.value for cell in new[1]]
        assert headers[:4] == ['Company', 'Tax', 'Source value', 'Website']
        assert headers[4:] == ['Merlin match', 'Merlin note', 'Website (Merlin)']
        assert new['D2'].value == 'customer website'
        assert new.cell(2, headers.index('Website (Merlin)') + 1).value == 'https://alfa.si/'
    finally:
        original.close(); exported.close()


def test_csv_uses_persisted_rows_and_preserves_duplicates_review_rows_and_scope(tmp_path):
    root, repository, service, workspace, upload, job, artifact = _run(
        tmp_path, ('EMAIL', 'PHONE'), source_format='CSV')
    stored_upload = repository.get_upload(upload.upload_id)
    # CSV export is deliberately based on the durable parsed rows, not these bytes.
    (root/stored_upload['relative_path']).write_text('changed after parsing', encoding='utf-8')
    exported_result = export_import_xlsx(repository, job.job_id, root)
    workbook = load_workbook(exported_result['path'], data_only=False)
    try:
        sheet = workbook['Merlin Results']
        headers = [cell.value for cell in sheet[1]]
        assert headers == ['Company', 'Tax', 'Source value', 'Merlin match',
                           'Merlin note', 'Email', 'Phone']
        assert [[sheet.cell(row, column).value for column in range(1, 4)]
                for row in range(2, 6)] == [
                    ['ALFA d.o.o.', '11111111', '=1+1'],
                    ['ALFA d.o.o.', '11111111', '=1+1'],
                    ['DVOJNIK d.o.o.', None, 'ambiguous'],
                    ['NEZNANO d.o.o.', None, 'not found']]
        assert sheet['C2'].data_type == 's'
        match_column = headers.index('Merlin match') + 1
        note_column = headers.index('Merlin note') + 1
        assert [sheet.cell(row, match_column).value for row in range(2, 6)] == [
            'Matched', 'Matched', 'Ambiguous', 'Not found']
        assert sheet.cell(2, note_column).value is None
        assert sheet.cell(4, note_column).value == (
            'Multiple possible company matches. Please review this row.')
        assert sheet.cell(5, note_column).value == 'No matching company was found.'
        visible_text = ' '.join(str(cell.value) for row in workbook['Merlin review']
                                for cell in row if cell.value is not None)
        assert job.job_id not in visible_text
        assert not any(term in visible_text.casefold() for term in (
            'sqlite', 'provider', 'adapter', 'run id', 'worker', 'traceback'))
    finally:
        workbook.close()

    current = service.artifacts(workspace, job.job_id)
    other = service.create_workspace('Other')['workspace_id']
    assert len(current) == 1 and current[0].artifact_id == artifact['artifact_id']
    assert service.artifacts(other, job.job_id) == ()
