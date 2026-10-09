"""Functional, no-network tests for the isolated MERLIN web application."""
import hashlib
import io
import subprocess
import sys
from pathlib import Path

import pytest
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from werkzeug.datastructures import FileStorage

from merlin.app import create_app


def _csrf(client):
    with client.session_transaction() as state:
        return state['csrf_token']


def _csv(name='companies.csv', rows=None):
    rows = rows or ('ALFA d.o.o.,11111111\n',)
    content = ('Company,Tax\n' + ''.join(rows)).encode()
    return io.BytesIO(content), name


def _xlsx(name='companies.xlsx'):
    stream = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['Company', 'Tax'])
    sheet.append(['ALFA d.o.o.', '11111111'])
    workbook.save(stream); workbook.close(); stream.seek(0)
    return stream, name


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    monkeypatch.setattr('requests.sessions.Session.request',
                        lambda *_args, **_kwargs: pytest.fail('network/provider call'))


@pytest.fixture
def web(tmp_path):
    app = create_app({
        'TESTING': True,
        'SECRET_KEY': 'test-only-secret-key',
        'MERLIN_DB': tmp_path/'control.sqlite3',
        'MERLIN_STORAGE_ROOT': tmp_path/'storage',
        'MERLIN_WORKSPACE_ID': 'workspace-a',
        'MERLIN_WORKSPACE_NAME': 'Workspace A',
        'UPLOAD_MAX_BYTES': 1024*1024,
        'UPLOAD_MAX_ROWS': 20,
        'SESSION_COOKIE_SECURE': False,
    })
    client = app.test_client()
    client.get('/')
    return app, client, app.extensions['merlin_service'], tmp_path/'storage'


def _upload(client, source=None):
    source = source or _csv()
    response = client.post('/uploads', data={
        '_csrf': _csrf(client), 'company_file': source,
    }, content_type='multipart/form-data')
    assert response.status_code == 303
    return response.headers['Location'].rsplit('/', 1)[-1]


def _queue(client, upload_id, outputs=('WEBSITE', 'EMAIL', 'PHONE', 'FINANCIALS'),
           extra=None):
    data = {'_csrf': _csrf(client), 'requested_outputs': list(outputs)}
    data.update(extra or {})
    response = client.post(f'/uploads/{upload_id}/enrich', data=data)
    assert response.status_code == 303
    return response.headers['Location'].rsplit('/', 1)[-1]


def _start(repository, job_id, worker='test-worker'):
    claimed = repository.claim_oldest(worker)
    assert claimed['job_id'] == job_id
    repository.transition(job_id, 'RUNNING', worker_id=worker)


def test_landing_is_isolated_accessible_and_security_hardened(web):
    app, client, _, _ = web
    response = client.get('/')
    html = response.get_data(as_text=True)
    assert response.status_code == 200
    assert 'Market Exploration, Research &amp; Lead Intelligence' in html
    assert 'Control Room' not in html and 'Monitor' not in html
    assert 'Powered by Dastabase' in html
    soup = BeautifulSoup(html, 'html.parser')
    assert soup.select_one('body > header') is None
    assert soup.select_one('.hero h1').get_text(strip=True) == 'MERLIN'
    assert soup.select_one('.hero .mark').get_text(strip=True) == '✦'
    file_input = soup.select_one('input[name="company_file"]')
    target = soup.select_one('[data-drop-area]')
    continue_button = soup.select_one('[data-upload-continue]')
    assert file_input['class'] == ['file-input']
    assert target['for'] == file_input['id']
    assert continue_button.has_attr('disabled') and continue_button.has_attr('hidden')
    assert 'Content-Security-Policy' in response.headers
    assert response.headers['Cache-Control'] == 'no-store'
    assert app.config['SESSION_COOKIE_HTTPONLY'] is True
    assert app.config['SESSION_COOKIE_SAMESITE'] == 'Lax'
    for path in ('/jobs', '/api/current-job', '/data', '/enrich'):
        assert client.get(path).status_code == 404


def test_merlin_app_imports_no_control_room_or_monitor_web_surface():
    command = ('import sys; import merlin.app; '
               "assert 'control_room.app' not in sys.modules; "
               "assert 'monitor.app' not in sys.modules; "
               "assert 'monitor.metrics' not in sys.modules")
    result = subprocess.run([sys.executable, '-c', command], cwd=Path.cwd(),
                            capture_output=True, text=True, check=False)
    assert result.returncode == 0, result.stderr


@pytest.mark.parametrize('source, expected_format', [
    (_csv(), 'CSV'), (_xlsx(), 'XLSX'),
])
def test_csv_and_xlsx_upload_reach_automatic_configuration(web, source, expected_format):
    _, client, service, _ = web
    upload_id = _upload(client, source)
    response = client.get(f'/uploads/{upload_id}')
    html = response.get_data(as_text=True)
    upload = service.upload('workspace-a', upload_id)
    assert response.status_code == 200
    assert upload.format == expected_format and upload.row_count == 1
    assert '1 rows detected' in html
    assert 'Company name' in html and 'Company' in html
    assert 'Which columns identify the company?' not in html
    soup = BeautifulSoup(html, 'html.parser')
    options = soup.select('input[name="requested_outputs"]')
    assert [item['value'] for item in options] == [
        'WEBSITE', 'EMAIL', 'PHONE', 'FINANCIALS']
    assert all(item.has_attr('checked') for item in options)


def test_invalid_upload_and_csrf_fail_safely(web):
    _, client, _, _ = web
    no_csrf = client.post('/uploads', data={
        'company_file': (io.BytesIO(b'bad'), 'bad.txt')},
        content_type='multipart/form-data')
    assert no_csrf.status_code == 400
    invalid = client.post('/uploads', data={
        '_csrf': _csrf(client), 'company_file': (io.BytesIO(b'bad'), 'bad.txt')},
        content_type='multipart/form-data')
    assert invalid.status_code == 400
    body = invalid.get_data(as_text=True)
    assert 'Only .csv and .xlsx files are accepted' in body
    assert 'Traceback' not in body and 'sqlite' not in body.lower()


def test_minimal_mapping_correction_only_when_required(web):
    _, client, _, _ = web
    upload_id = _upload(client, (
        io.BytesIO(b'Unknown,Other\nALFA d.o.o.,value\n'), 'unknown.csv'))
    html = client.get(f'/uploads/{upload_id}').get_data(as_text=True)
    assert 'Which columns identify the company?' in html
    job_id = _queue(client, upload_id, ('WEBSITE',), {'mapping_company_name': '0'})
    assert client.get(f'/jobs/{job_id}').status_code == 200


def test_configuration_requires_outputs_and_queues_owned_import_job(web):
    _, client, service, _ = web
    upload_id = _upload(client)
    empty = client.post(f'/uploads/{upload_id}/enrich', data={'_csrf': _csrf(client)})
    assert empty.status_code == 400
    assert 'At least one requested output is required' in empty.get_data(as_text=True)

    job_id = _queue(client, upload_id, ('WEBSITE', 'FINANCIALS'))
    stored = service.repository.get_job(job_id)
    assert stored['module'] == 'IMPORT_ENRICH'
    assert stored['execution_adapter'] == 'IMPORT_ENRICH'
    assert stored['workspace_id'] == 'workspace-a'
    assert stored['origin_surface'] == 'MERLIN'
    assert stored['requested_outputs_json'] == '["WEBSITE","FINANCIALS"]'


def test_unknown_and_cross_workspace_resources_are_safe_404(web):
    _, client, service, _ = web
    workspace_b = service.create_workspace('Workspace B', 'workspace-b')['workspace_id']
    other_upload = service.ingest_upload(workspace_b,
        FileStorage(stream=_csv()[0], filename='other.csv', content_type='text/csv'))
    other_job = service.queue_import(workspace_b, other_upload.upload_id, 'Other',
        {'company_name': 0, 'tax_number': 1}, ('WEBSITE',))

    for path in (f'/uploads/{other_upload.upload_id}', '/uploads/unknown',
                 f'/jobs/{other_job.job_id}', '/jobs/unknown',
                 f'/api/jobs/{other_job.job_id}', '/api/jobs/unknown',
                 f'/jobs/{other_job.job_id}/download', '/jobs/unknown/download'):
        assert client.get(path).status_code == 404


def test_polling_api_is_safe_and_reflects_persisted_progress(web):
    _, client, service, _ = web
    upload_id = _upload(client, _csv(rows=(
        'ALFA d.o.o.,11111111\n', 'BETA d.o.o.,22222222\n')))
    job_id = _queue(client, upload_id, ('EMAIL',))
    queued = client.get(f'/api/jobs/{job_id}').get_json()['job']
    assert set(queued) == {'state', 'stage', 'processed', 'total', 'percent',
        'download_ready', 'requested_outputs', 'matched_rows', 'review_rows'}
    assert queued == {**queued, 'state': 'working', 'stage': 'Identifying companies',
                      'processed': 0, 'total': 2, 'percent': 0,
                      'download_ready': False, 'requested_outputs': ['EMAIL']}
    page = BeautifulSoup(client.get(f'/jobs/{job_id}').get_data(as_text=True),
                         'html.parser')
    progress = page.select_one('[data-job-progress]')
    working = page.select_one('[data-working-status]')
    assert progress['data-state'] == 'working'
    assert len(page.select('.sparkles span')) == 3
    assert not working.has_attr('hidden')
    assert working.select_one('p').get_text(strip=True) == 'Working on it…'
    assert page.select_one('[data-eta-fallback]').get_text(' ', strip=True) == (
        'Time estimate will appear once Merlin has enough information.')
    assert not any(key in queued for key in ('eta', 'eta_seconds', 'remaining_seconds'))
    css = client.get('/static/merlin.css').get_data(as_text=True)
    assert '@media (prefers-reduced-motion: reduce)' in css

    repository = service.repository
    _start(repository, job_id)
    repository.set_import_stage(job_id, 'EXISTING_ENRICHMENT', 'test-worker')
    checking = client.get(f'/api/jobs/{job_id}').get_json()['job']
    assert checking['stage'] == 'Checking existing information'
    repository.set_import_stage(job_id, 'DISCOVERY', 'test-worker')
    repository.update_progress(job_id, 1, 0, 0, 0, 'test-worker')
    active = client.get(f'/api/jobs/{job_id}').get_json()['job']
    assert (active['stage'], active['processed'], active['total'], active['percent']) == (
        'Finding missing information', 1, 2, 50)
    repository.set_import_stage(job_id, 'EXPORT', 'test-worker')
    preparing = client.get(f'/api/jobs/{job_id}').get_json()['job']
    assert preparing['stage'] == 'Preparing your Excel file'
    serialized = str(active).lower()
    assert not any(term in serialized for term in (
        'worker', 'adapter', 'run_id', 'provider', 'path', 'error_message'))


def test_completed_job_downloads_only_verified_final_workspace_artifact(web):
    _, client, service, root = web
    upload_id = _upload(client)
    job_id = _queue(client, upload_id, ('WEBSITE',))
    repository = service.repository
    _start(repository, job_id)

    directory = root/'jobs'/job_id
    directory.mkdir(parents=True)
    internal = directory/'manifest.txt'; internal.write_text('internal')
    repository.add_artifact(job_id, 'MANIFEST', f'jobs/{job_id}/manifest.txt',
                            internal.stat().st_size,
                            hashlib.sha256(internal.read_bytes()).hexdigest())
    repository.transition(job_id, 'COMPLETED', worker_id='test-worker')
    assert client.get(f'/jobs/{job_id}/download').status_code == 404

    final = directory/'final.xlsx'
    workbook = Workbook(); workbook.active['A1'] = 'MERLIN'; workbook.save(final); workbook.close()
    repository.add_artifact(job_id, 'UPLOAD_RECONCILIATION',
                            f'jobs/{job_id}/final.xlsx', final.stat().st_size,
                            hashlib.sha256(final.read_bytes()).hexdigest())

    progress = client.get(f'/api/jobs/{job_id}').get_json()['job']
    assert progress['state'] == 'done' and progress['download_ready'] is True
    page = BeautifulSoup(client.get(f'/jobs/{job_id}').get_data(as_text=True),
                         'html.parser')
    assert page.select_one('[data-job-progress]')['data-state'] == 'done'
    assert page.select_one('[data-working-status]').has_attr('hidden')
    assert not page.select_one('[data-complete]').has_attr('hidden')
    start_again = page.find('a', string='Enrich another file')
    assert start_again is not None and start_again['href'] == '/'
    response = client.get(f'/jobs/{job_id}/download')
    assert response.status_code == 200
    assert response.headers['Content-Type'].startswith(
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    assert 'attachment' in response.headers['Content-Disposition']
    downloaded = load_workbook(io.BytesIO(response.data))
    assert downloaded.active['A1'].value == 'MERLIN'; downloaded.close()
    final.write_bytes(b'tampered')
    assert client.get(f'/jobs/{job_id}/download').status_code == 404
    assert client.get(f'/api/jobs/{job_id}').get_json()['job']['state'] == 'failed'


def test_failed_job_never_exposes_internal_error(web):
    _, client, service, _ = web
    upload_id = _upload(client)
    job_id = _queue(client, upload_id, ('WEBSITE',))
    repository = service.repository
    _start(repository, job_id)
    secret = 'OperationalError sqlite /secret/path provider payload Traceback'
    repository.transition(job_id, 'FAILED', worker_id='test-worker',
                          error_code='IMPORT_ENRICH_ADAPTER_ERROR', error_message=secret)
    api = client.get(f'/api/jobs/{job_id}')
    page = client.get(f'/jobs/{job_id}')
    assert api.get_json()['job']['state'] == 'failed'
    rendered = BeautifulSoup(page.get_data(as_text=True), 'html.parser')
    assert rendered.select_one('[data-job-progress]')['data-state'] == 'failed'
    assert rendered.select_one('[data-working-status]').has_attr('hidden')
    combined = api.get_data(as_text=True) + page.get_data(as_text=True)
    assert secret not in combined
    assert '/secret/path' not in combined and 'OperationalError' not in combined
