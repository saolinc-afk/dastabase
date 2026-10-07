"""Lossless XLSX result export tests for Import & Enrich."""
import hashlib
import json
import sqlite3

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from control_room.import_enrich_adapter import ImportEnrichAdapter
from control_room.import_export import EXPORT_HEADERS, export_import_xlsx
from control_room.repository import JobRepository
from control_room.worker import Worker


def _digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _canonical(path):
    conn = sqlite3.connect(path)
    conn.executescript('''
      CREATE TABLE companies_lite(
        id INTEGER PRIMARY KEY, company_name TEXT, registration_number TEXT,
        tax_number TEXT, address TEXT, municipality TEXT, revenue_2025 REAL,
        profit_2025 REAL, employees_2025 REAL);
      INSERT INTO companies_lite VALUES
        (1,'ALFA d.o.o.','1001','11111111','Alfa 1','Kranj',100,12,3),
        (2,'DVOJNIK d.o.o.','1002','22222222','Prva 2','Celje',200,22,4),
        (3,'DVOJNIK d.o.o.','1003','33333333','Druga 3','Maribor',300,32,5),
        (4,'BETA d.o.o.','1004','44444444','Beta 4','Koper',444,44,6);
      CREATE TABLE website_discovery(id INTEGER PRIMARY KEY,company_id INTEGER,
        website TEXT,status TEXT,verified_scope TEXT,relationship TEXT);
      INSERT INTO website_discovery VALUES
        (1,1,'https://alfa.si/','VERIFIED','https://alfa.si/','LEGAL_ENTITY'),
        (2,4,'https://beta.si/','VERIFIED','https://beta.si/','LEGAL_ENTITY');
      CREATE TABLE email_discovery(id INTEGER PRIMARY KEY,company_id INTEGER,
        email TEXT,website TEXT);
      INSERT INTO email_discovery VALUES
        (1,1,'info@alfa.si','https://alfa.si/'),
        (2,4,'hello@beta.si','https://beta.si/');
    ''')
    conn.commit()
    conn.close()


def _workbook(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Prijave'
    sheet.freeze_panes = 'B2'
    sheet.append(['Ime', 'E-pošta', 'Podjetje', 'Davčna', 'Izračun', 'Website'])
    sheet.append(['Ana', 'ana@alfa.si', 'ALFA d.o.o.', '11111111', '=1+1', 'source-value'])
    sheet.append(['Ana', 'ana@alfa.si', 'ALFA d.o.o.', '11111111', '=1+1', 'source-value'])
    sheet.append(['Eva', '', 'DVOJNIK d.o.o.', '', '=3+3', 'ambiguous'])
    sheet.append(['', '', '', '', '', ''])
    sheet.append(['Žan', '', 'NEZNANO podjetje', '', '=4+4', 'unresolved'])
    sheet['A1'].font = Font(bold=True, color='FFFFFF')
    sheet['A1'].fill = PatternFill('solid', fgColor='204060')
    sheet.row_dimensions[2].height = 27
    sheet.column_dimensions['A'].width = 24
    hidden = workbook.create_sheet('Skriti podatki')
    hidden.sheet_state = 'hidden'
    hidden['A1'] = 'ostane'
    notes = workbook.create_sheet('Opombe')
    notes.merge_cells('A1:C1')
    notes['A1'] = 'Združeno'
    # Force collision-safe QC sheet naming without allowing re-export duplication.
    workbook.create_sheet('Dastabase QC')['A1'] = 'Izvorna stran'
    workbook.save(path)


def _setup(tmp_path):
    root = tmp_path / 'storage'
    source_dir = root / 'uploads'
    source_dir.mkdir(parents=True)
    source = source_dir / 'u.xlsx'
    _workbook(source)
    canonical = tmp_path / 'canonical.sqlite3'
    _canonical(canonical)
    repository = JobRepository(tmp_path / 'control.sqlite3')
    repository.initialize()
    rows = [
        ['Ana', 'ana@alfa.si', 'ALFA d.o.o.', '11111111', '=1+1', 'source-value'],
        ['Ana', 'ana@alfa.si', 'ALFA d.o.o.', '11111111', '=1+1', 'source-value'],
        ['Eva', '', 'DVOJNIK d.o.o.', '', '=3+3', 'ambiguous'],
        ['', '', '', '', '', ''],
        ['Žan', '', 'NEZNANO podjetje', '', '=4+4', 'unresolved'],
    ]
    repository.create_upload({
        'upload_id': 'u', 'original_filename': 'prijave.xlsx',
        'relative_path': 'uploads/u.xlsx', 'sha256': _digest(source),
        'size_bytes': source.stat().st_size, 'format': 'XLSX',
        'worksheet_name': 'Prijave',
        'headers': ['Ime', 'E-pošta', 'Podjetje', 'Davčna', 'Izračun', 'Website'],
        'rows': rows,
    })
    mapping = {'person_name': 0, 'email': 1, 'company_name': 2, 'tax_number': 3}
    job = repository.create_import_enrich_job('u', 'Izvoz prijav', mapping)
    return root, source, canonical, repository, job


def _update_upload_source(repository, source, rows=None, relative_path=None, format='XLSX'):
    conn = sqlite3.connect(repository.path)
    if rows is not None:
        for row_number, values in enumerate(rows, 2):
            conn.execute('UPDATE upload_rows SET original_values_json=? '
                         'WHERE upload_id=? AND row_number=?',
                         (json.dumps(values, ensure_ascii=False), 'u', row_number))
    conn.execute('UPDATE uploads SET sha256=?,size_bytes=?,relative_path=?,format=? '
                 'WHERE upload_id=?', (_digest(source), source.stat().st_size,
                 relative_path or 'uploads/u.xlsx', format, 'u'))
    conn.commit()
    conn.close()


def test_import_export_preserves_workbook_and_maps_every_durable_source_row(tmp_path):
    root, source, canonical, repository, job = _setup(tmp_path)
    source_hash = _digest(source)
    adapter = ImportEnrichAdapter(repository, canonical, storage_root=root)
    worker = Worker(repository, adapter, worker_id='export-worker',
                    adapters={'IMPORT_ENRICH': adapter})
    assert worker.run_once()
    completed = repository.get_job(job['job_id'])
    assert completed['status'] == 'COMPLETED'
    assert completed['progress_stage'] == 'EXPORT'

    artifact = repository.artifacts(job['job_id'])
    assert len(artifact) == 1
    assert artifact[0]['artifact_type'] == 'UPLOAD_RECONCILIATION'
    output = root / artifact[0]['relative_path']
    assert output.is_file() and _digest(output) == artifact[0]['sha256']
    assert _digest(source) == source_hash

    original = load_workbook(source, data_only=False)
    exported = load_workbook(output, data_only=False)
    try:
        assert exported.sheetnames == [
            'Prijave', 'Skriti podatki', 'Opombe', 'Dastabase QC', 'Dastabase QC (2)']
        old, new = original['Prijave'], exported['Prijave']
        assert new.freeze_panes == old.freeze_panes == 'B2'
        assert new['E2'].value == '=1+1'
        assert new['A1'].font.bold == old['A1'].font.bold is True
        assert new['A1'].font.color.rgb == old['A1'].font.color.rgb == '00FFFFFF'
        assert new['A1'].fill.fgColor.rgb == old['A1'].fill.fgColor.rgb == '00204060'
        assert new.row_dimensions[2].height == 27
        assert new.column_dimensions['A'].width == 24
        assert exported['Skriti podatki'].sheet_state == 'hidden'
        assert list(exported['Opombe'].merged_cells.ranges)[0].coord == 'A1:C1'
        assert exported['Dastabase QC']['A1'].value == 'Izvorna stran'

        headers = {new.cell(1, column).value: column
                   for column in range(1, new.max_column + 1)}
        assert 'Website (Dastabase)' in headers
        row_id = headers['Dastabase Row ID']
        assert [new.cell(row, row_id).value for row in range(2, 7)] == [
            f'{job["job_id"]}:1', f'{job["job_id"]}:2', f'{job["job_id"]}:3',
            f'{job["job_id"]}:4', f'{job["job_id"]}:5']
        status = headers['Match Status']
        assert [new.cell(row, status).value for row in range(2, 7)] == [
            'MATCHED', 'MATCHED', 'AMBIGUOUS', 'UNRESOLVED', 'UNRESOLVED']
        company = headers['Canonical Company ID']
        website = headers['Website (Dastabase)']
        email = headers['Default Email']
        assert new.cell(2, company).value == new.cell(3, company).value == 1
        assert new.cell(2, website).value == 'https://alfa.si/'
        assert new.cell(2, email).value == 'info@alfa.si'
        assert new.cell(4, company).value is None
        assert new.cell(6, website).value is None
        qc = {row[0].value: row[1].value for row in exported['Dastabase QC (2)'].iter_rows()
              if row[0].value}
        assert qc['Source registration rows'] == qc['Exported registration rows'] == 5
        assert (qc['Matched rows'], qc['Ambiguous rows'], qc['Unresolved rows']) == (2, 1, 2)
        assert qc['Duplicate source row groups'] == 1
        assert qc['Duplicate source rows preserved'] == 1
        assert qc['Rows with website'] == qc['Rows with default email'] == 2
        assert qc['Rows with financials'] == 2
    finally:
        original.close()
        exported.close()


def test_reexport_rebuilds_from_source_without_duplicate_columns_or_qc_sheets(tmp_path):
    root, source, canonical, repository, job = _setup(tmp_path)
    adapter = ImportEnrichAdapter(repository, canonical, storage_root=root)
    Worker(repository, adapter, worker_id='export-worker',
           adapters={'IMPORT_ENRICH': adapter}).run_once()
    first = repository.artifacts(job['job_id'])[0]
    result = export_import_xlsx(repository, job['job_id'], root)
    second = repository.artifacts(job['job_id'])[0]
    assert first['artifact_id'] == second['artifact_id']
    assert first['relative_path'] != second['relative_path']
    assert (root / first['relative_path']).is_file()
    assert result['qc']['Exported registration rows'] == 5
    workbook = load_workbook(result['path'], data_only=False)
    try:
        headers = [cell.value for cell in workbook['Prijave'][1]]
        assert sum(value.startswith('Dastabase Row ID') for value in headers if value) == 1
        assert workbook.sheetnames.count('Dastabase QC (2)') == 1
        assert all(any(value == header or str(value).startswith(header + ' (Dastabase')
                       for value in headers) for header in EXPORT_HEADERS)
    finally:
        workbook.close()


def test_export_rejects_changed_source_workbook(tmp_path):
    root, source, canonical, repository, job = _setup(tmp_path)
    adapter = ImportEnrichAdapter(repository, canonical)
    Worker(repository, adapter, worker_id='match-worker',
           adapters={'IMPORT_ENRICH': adapter}).run_once()
    source.write_bytes(source.read_bytes() + b'changed')
    try:
        export_import_xlsx(repository, job['job_id'], root)
    except ValueError as exc:
        assert 'hash does not match' in str(exc)
    else:
        raise AssertionError('changed source workbook was accepted')


def test_failed_registration_preserves_previous_artifact_and_immutable_file(tmp_path, monkeypatch):
    root, source, canonical, repository, job = _setup(tmp_path)
    adapter = ImportEnrichAdapter(repository, canonical, storage_root=root)
    Worker(repository, adapter, worker_id='export-worker',
           adapters={'IMPORT_ENRICH': adapter}).run_once()
    previous = repository.artifacts(job['job_id'])[0]
    previous_path = root / previous['relative_path']
    previous_bytes = previous_path.read_bytes()

    def fail_registration(*args, **kwargs):
        raise RuntimeError('injected failure before artifact registration')

    monkeypatch.setattr(repository, 'add_artifact', fail_registration)
    with pytest.raises(RuntimeError, match='injected failure'):
        export_import_xlsx(repository, job['job_id'], root)

    registered = repository.artifacts(job['job_id'])[0]
    assert registered == previous
    assert previous_path.read_bytes() == previous_bytes
    assert previous_path.stat().st_size == previous['size_bytes']
    assert _digest(previous_path) == previous['sha256']
    published = list((root / 'jobs' / job['job_id']).glob('import-enriched-*.xlsx'))
    assert len(published) == 2
    load_workbook(previous_path).close()
    assert not list((root / 'jobs' / job['job_id']).glob('.import-enriched-*.xlsx'))


def test_adjacent_matched_rows_keep_distinct_company_data(tmp_path):
    root, source, canonical, repository, job = _setup(tmp_path)
    workbook = load_workbook(source)
    workbook['Prijave'].cell(3, 1, 'Bina')
    workbook['Prijave'].cell(3, 2, 'hello@beta.si')
    workbook['Prijave'].cell(3, 3, 'BETA d.o.o.')
    workbook['Prijave'].cell(3, 4, '44444444')
    workbook.save(source)
    workbook.close()
    rows = repository.upload_rows('u')
    values = [row['original_values'] for row in rows]
    values[1][:4] = ['Bina', 'hello@beta.si', 'BETA d.o.o.', '44444444']
    _update_upload_source(repository, source, values)

    adapter = ImportEnrichAdapter(repository, canonical, storage_root=root)
    Worker(repository, adapter, worker_id='export-worker',
           adapters={'IMPORT_ENRICH': adapter}).run_once()
    output = root / repository.artifacts(job['job_id'])[0]['relative_path']
    exported = load_workbook(output, data_only=False)
    try:
        sheet = exported['Prijave']
        headers = {cell.value: cell.column for cell in sheet[1]}
        expected = {
            2: (1, 'ALFA d.o.o.', 'https://alfa.si/', 'info@alfa.si', 100, 12),
            3: (4, 'BETA d.o.o.', 'https://beta.si/', 'hello@beta.si', 444, 44),
        }
        for row_number, values in expected.items():
            assert tuple(sheet.cell(row_number, headers[name]).value for name in (
                'Canonical Company ID', 'Matched Company', 'Website (Dastabase)',
                'Default Email', 'Revenue 2025', 'Profit 2025')) == values
    finally:
        exported.close()


def test_persisted_non_first_source_worksheet_is_the_only_enriched_sheet(tmp_path):
    root, source, canonical, repository, job = _setup(tmp_path)
    workbook = load_workbook(source)
    cover = workbook.create_sheet('Naslovnica', 0)
    cover['A1'] = 'Nespremenjena naslovnica'
    cover['B2'] = '=40+2'
    workbook.save(source)
    workbook.close()
    _update_upload_source(repository, source)

    adapter = ImportEnrichAdapter(repository, canonical, storage_root=root)
    Worker(repository, adapter, worker_id='export-worker',
           adapters={'IMPORT_ENRICH': adapter}).run_once()
    exported = load_workbook(root / repository.artifacts(job['job_id'])[0]['relative_path'])
    try:
        assert exported.sheetnames[0] == 'Naslovnica'
        assert exported['Naslovnica']['A1'].value == 'Nespremenjena naslovnica'
        assert exported['Naslovnica']['B2'].value == '=40+2'
        assert exported['Naslovnica'].max_column == 2
        assert any(cell.value == 'Dastabase Row ID' for cell in exported['Prijave'][1])
    finally:
        exported.close()


@pytest.mark.parametrize('replacement', (1, 2, 999))
def test_invalid_duplicate_or_out_of_range_source_row_fails_closed(tmp_path, replacement):
    root, source, canonical, repository, job = _setup(tmp_path)
    adapter = ImportEnrichAdapter(repository, canonical)
    Worker(repository, adapter, worker_id='match-worker',
           adapters={'IMPORT_ENRICH': adapter}).run_once()
    conn = sqlite3.connect(repository.path)
    conn.execute('PRAGMA foreign_keys=OFF')
    conn.execute('UPDATE job_items SET upload_row_number=? WHERE job_id=? AND item_position=2',
                 (replacement, job['job_id']))
    conn.commit()
    conn.close()
    with pytest.raises(ValueError, match='invalid, duplicate, or out-of-range'):
        export_import_xlsx(repository, job['job_id'], root)
    assert repository.artifacts(job['job_id']) == []


def test_csv_job_with_storage_root_does_not_attempt_xlsx_export(tmp_path):
    root, source, canonical, repository, job = _setup(tmp_path)
    csv_source = root / 'uploads' / 'u.csv'
    csv_source.write_text('Ime,E-pošta,Podjetje,Davčna\nAna,ana@alfa.si,ALFA d.o.o.,11111111\n',
                          encoding='utf-8')
    _update_upload_source(repository, csv_source, relative_path='uploads/u.csv', format='CSV')
    adapter = ImportEnrichAdapter(repository, canonical, storage_root=root)
    Worker(repository, adapter, worker_id='csv-worker',
           adapters={'IMPORT_ENRICH': adapter}).run_once()
    assert repository.get_job(job['job_id'])['status'] == 'COMPLETED'
    assert repository.get_job(job['job_id'])['progress_stage'] == 'EXISTING_ENRICHMENT'
    assert repository.artifacts(job['job_id']) == []


def test_identity_states_and_scope_are_exported_without_invented_company_facts(tmp_path):
    root, source, canonical, repository, job = _setup(tmp_path)
    adapter = ImportEnrichAdapter(repository, canonical)
    Worker(repository, adapter, worker_id='match-worker',
           adapters={'IMPORT_ENRICH': adapter}).run_once()
    items = repository.job_items(job['job_id'])
    external = items[4]
    assert external['match_status'] == 'UNRESOLVED' and external['identity_task_id']
    conn = sqlite3.connect(repository.path)
    conn.execute("UPDATE job_items SET identity_status='RESOLVED_NEW_ENTITY' "
                 'WHERE job_id=? AND item_position=?', (job['job_id'], external['item_position']))
    conn.execute("UPDATE identity_resolution_tasks SET "
                 "identity_resolution_status='RESOLVED_NEW_ENTITY',"
                 "enrichment_scope='BULK_IN_SCOPE' WHERE task_id=?",
                 (external['identity_task_id'],))
    conn.commit()
    conn.close()

    result = export_import_xlsx(repository, job['job_id'], root)
    exported = load_workbook(result['path'])
    try:
        sheet = exported['Prijave']
        headers = {cell.value: cell.column for cell in sheet[1]}
        assert sheet.cell(6, headers['Identity Status']).value == 'RESOLVED_NEW_ENTITY'
        assert sheet.cell(6, headers['Research Scope']).value == 'BULK_IN_SCOPE'
        for header in ('Canonical Company ID', 'Matched Company', 'Website (Dastabase)',
                       'Default Email', 'Revenue 2025', 'Tax Number'):
            assert sheet.cell(6, headers[header]).value is None
        assert sheet.cell(5, headers['Identity Status']).value == 'NOT_ELIGIBLE'
        assert result['qc']['Externally identified new entities'] == 1
        assert result['qc']['Skipped / not eligible rows'] >= 1
    finally:
        exported.close()
