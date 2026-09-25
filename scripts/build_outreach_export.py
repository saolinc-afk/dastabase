"""Read-only outreach export under current Phase 2 attribution semantics."""
import csv
import json
import re
import sqlite3
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from discovery.ownership import email_attribution

DB = ROOT / 'database/dastabase_lite.db'
OUT = ROOT / 'exports'
OUT.mkdir(exist_ok=True)
STAMP = datetime.now().strftime('%Y%m%d_%H%M%S')
BASE = OUT / f'outreach_gorenjska_goriska_{STAMP}'
HEADERS = [
    'dastabase_id', 'region', 'company_name', 'address', 'postal_code', 'city', 'municipality',
    'tax_number', 'registration_number', 'revenue_2025', 'employees_2025', 'website',
    'website_status', 'website_confidence', 'primary_email', 'other_emails', 'email_count',
    'contact_person', 'sent', 'sent_date', 'reply', 'called', 'call_date', 'outcome',
    'follow_up_date', 'owner', 'notes',
]


def obj(raw, fallback):
    try:
        value = json.loads(raw or fallback)
        return value if isinstance(value, type(fallback)) else fallback
    except (TypeError, json.JSONDecodeError):
        return fallback


def postal_parts(value):
    match = re.match(r'^\s*(\d{4})\s*(.*)$', value or '')
    return (match.group(1), match.group(2).strip()) if match else ('', '')


def numeric(value):
    if value in (None, ''):
        return None
    return float(value) if isinstance(value, str) else value


def primary_rank(email, website):
    local = email['email'].split('@', 1)[0].lower()
    generic = local in {'info', 'contact', 'kontakt', 'office', 'sales', 'prodaja', 'support', 'podpora'}
    host = re.sub(r'^www\.', '', re.sub(r'^https?://', '', website or '').split('/')[0].lower())
    same_host = email['email'].rsplit('@', 1)[-1].lower() == host
    return (-int(generic), -int(same_host), -int(email.get('confidence') or 0), email['email'])


def main():
    conn = sqlite3.connect(DB)
    conn.row_factory = sqlite3.Row
    integrity = conn.execute('PRAGMA integrity_check').fetchone()[0]
    foreign_keys = [tuple(row) for row in conn.execute('PRAGMA foreign_key_check')]
    companies = conn.execute('''SELECT c.*, w.id AS website_row_id, w.website, w.status AS website_status,
        w.confidence AS website_confidence, w.rule_version, w.evidence_json, w.ownership_json
        FROM companies_lite c LEFT JOIN website_discovery w ON w.id=(
          SELECT id FROM website_discovery WHERE company_id=c.id ORDER BY id DESC LIMIT 1)
        WHERE trim(c.municipality) GLOB '4[0-9][0-9][0-9]*'
           OR trim(c.municipality) GLOB '5[0-9][0-9][0-9]*'
        ORDER BY c.id''').fetchall()
    company_ids = [row['id'] for row in companies]
    placeholders = ','.join('?' for _ in company_ids)
    mails = conn.execute(f'''SELECT * FROM email_discovery WHERE company_id IN ({placeholders})
                             ORDER BY company_id, confidence DESC, email''', company_ids).fetchall()
    by_company = defaultdict(list)
    for mail in mails:
        by_company[mail['company_id']].append(dict(mail))

    rows, excluded = [], []
    for source in companies:
        company = dict(source)
        postal, city = postal_parts(company['municipality'])
        region = 'Gorenjska (4xxx)' if postal.startswith('4') else 'Goriška (5xxx)'
        ownership = obj(company['ownership_json'], {})
        pages = obj(company['evidence_json'], [])
        valid = []
        for mail in by_company[company['id']]:
            evidence = obj(mail.get('evidence_json'), {})
            # A current hardened VERIFIED website with its stored ownership
            # proof is the gate before any address can leave the database.
            if company['website_status'] != 'VERIFIED' or company['rule_version'] != 'phase2-ownership-1':
                excluded.append((company['id'], mail['email'], 'non-current or non-VERIFIED website status'))
                continue
            attribution = email_attribution(company, mail['email'], mail.get('page_url') or '',
                                            evidence.get('publication', ''), ownership, pages,
                                            mail.get('page_title') or '', evidence.get('contact_block'),
                                            evidence.get('visible_email', ''))
            if attribution.get('attributable'):
                valid.append(mail)
            else:
                excluded.append((company['id'], mail['email'], attribution.get('reason', 'insufficient provenance')))
        valid.sort(key=lambda email: primary_rank(email, company['website']))
        primary = valid[0]['email'] if valid else ''
        other = '; '.join(email['email'] for email in valid[1:])
        rows.append({
            'dastabase_id': str(company['id']), 'region': region, 'company_name': company['company_name'] or '',
            'address': company['address'] or '', 'postal_code': postal, 'city': city,
            'municipality': company['municipality'] or '', 'tax_number': str(company['tax_number'] or ''),
            'registration_number': str(company['registration_number'] or ''),
            'revenue_2025': numeric(company['revenue_2025']), 'employees_2025': numeric(company['employees_2025']),
            'website': company['website'] or '', 'website_status': company['website_status'] or '',
            'website_confidence': company['website_confidence'] if company['website_confidence'] is not None else '',
            'primary_email': primary, 'other_emails': other, 'email_count': len(valid),
            'contact_person': '', 'sent': '', 'sent_date': '', 'reply': '', 'called': '', 'call_date': '',
            'outcome': '', 'follow_up_date': '', 'owner': '', 'notes': '',
        })
    rows.sort(key=lambda row: (row['region'], row['postal_code'], row['city'], row['company_name']))
    with BASE.with_suffix('.csv').open('w', newline='', encoding='utf-8-sig') as handle:
        writer = csv.DictWriter(handle, fieldnames=HEADERS)
        writer.writeheader(); writer.writerows(rows)

    wb = Workbook()
    wb.remove(wb.active)
    for title, subset in (('Outreach', rows), ('Mailing', [row for row in rows if row['primary_email']]),
                          ('Calling', [row for row in rows if not row['primary_email']])):
        sheet = wb.create_sheet(title)
        sheet.append(HEADERS)
        for row in subset:
            sheet.append([row[key] for key in HEADERS])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = f'A1:{get_column_letter(len(HEADERS))}{max(1, len(subset)+1)}'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F4E78')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[1].height = 30
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for column in ('A', 'E', 'H', 'I'):
            for cell in sheet[column]:
                cell.number_format = '@'
        for column in ('J', 'K'):
            for cell in sheet[column][1:]:
                cell.number_format = '#,##0.00' if column == 'J' else '#,##0'
        widths = [13, 19, 34, 36, 12, 18, 22, 14, 18, 16, 13, 36, 17, 12, 31, 48, 11,
                  22, 10, 13, 18, 10, 13, 22, 16, 18, 42]
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
    wb.save(BASE.with_suffix('.xlsx'))

    counts = {}
    for region in ('Gorenjska (4xxx)', 'Goriška (5xxx)'):
        subset = [row for row in rows if row['region'] == region]
        counts[region] = {'total_companies': len(subset), 'with_primary_email': sum(bool(row['primary_email']) for row in subset),
                          'without_email': sum(not row['primary_email'] for row in subset),
                          'attributable_email_addresses': sum(row['email_count'] for row in subset)}
    primary_counts = Counter(row['primary_email'].lower() for row in rows if row['primary_email'])
    report = {'integrity_check': integrity, 'foreign_key_check': foreign_keys, 'regions': counts,
              'combined': {'total_rows': len(rows), 'with_primary_email': sum(bool(row['primary_email']) for row in rows),
                           'without_email': sum(not row['primary_email'] for row in rows),
                           'attributable_email_addresses': sum(row['email_count'] for row in rows),
                           'duplicate_company_ids': [key for key, count in Counter(row['dastabase_id'] for row in rows).items() if count > 1],
                           'duplicate_primary_email_addresses': {key: count for key, count in primary_counts.items() if count > 1},
                           'rows_with_email_despite_nonattributable_status': sum(bool(row['primary_email']) and row['website_status'] != 'VERIFIED' for row in rows),
                           'excluded_suspicious_or_stale_email_records': len(excluded),
                           'exclusion_reasons': dict(Counter(reason for _, _, reason in excluded))},
              'files': {'xlsx': str(BASE.with_suffix('.xlsx')), 'csv': str(BASE.with_suffix('.csv'))}}
    conn.close()
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
