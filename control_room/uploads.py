"""Bounded, deterministic CSV/XLSX parsing and safe upload storage."""
import csv
import hashlib
import io
import os
import re
import uuid
import zipfile
import zlib
from pathlib import Path

from openpyxl import load_workbook


FIELDS = ('company_name', 'tax_number', 'registration_number', 'address', 'municipality')
IMPORT_FIELDS = ('person_name', 'email', 'phone', *FIELDS)
# XLSX is a ZIP container. These limits bound work before openpyxl expands XML
# and shared-string content. They intentionally sit well above normal Merlin
# uploads while rejecting compressed archives with disproportionate expansion.
XLSX_MAX_ZIP_ENTRIES = 1_000
XLSX_MAX_UNCOMPRESSED_BYTES = 128 * 1024 * 1024
XLSX_MAX_ENTRY_BYTES = 64 * 1024 * 1024
XLSX_MAX_COMPRESSION_RATIO = 200
XLSX_COMPRESSION_RATIO_MIN_BYTES = 1024 * 1024
HEADER_ALIASES = {
    'person_name': {'ime', 'name', 'ime in priimek', 'full name'},
    'email': {'email', 'e mail', 'e posta', 'elektronska posta'},
    'phone': {'telefon', 'phone', 'mobile', 'mobitel', 'mobilni telefon'},
    'company_name': {'naziv', 'naziv podjetja', 'podjetje', 'firma', 'company', 'company name'},
    'tax_number': {'davcna', 'davcna stevilka', 'davcna st', 'tax number', 'vat number'},
    'registration_number': {'maticna', 'maticna stevilka', 'maticna st', 'registration number'},
    'address': {'naslov', 'address', 'ulica'},
    'municipality': {'obcina', 'posta', 'municipality', 'city', 'kraj'},
}


class UploadError(ValueError):
    pass


def _plain(value):
    return '' if value is None else str(value).strip()


def header_key(value):
    import unicodedata
    text = unicodedata.normalize('NFKD', _plain(value)).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', ' ', text).strip()


def _suggest_mapping(headers, fields):
    suggestions = {}
    for field in fields:
        aliases = HEADER_ALIASES[field]
        matches = [index for index, header in enumerate(headers) if header_key(header) in aliases]
        suggestions[field] = matches[0] if len(matches) == 1 else None
    return suggestions


def suggest_mapping(headers):
    """Preserve the existing company-list mapping contract."""
    return _suggest_mapping(headers, FIELDS)


def suggest_import_mapping(headers):
    return _suggest_mapping(headers, IMPORT_FIELDS)


def normalize_import_row(headers, values, mapping=None, row_key=None):
    """Create a matching-only registration view without altering source values."""
    from control_room.matching import (normalize_email, normalize_name, normalize_phone,
                                       normalize_registration, normalize_tax, normalize_text)
    if len(headers) != len(values):
        raise UploadError('Registration row width does not match its headers')
    mapping = suggest_import_mapping(headers) if mapping is None else dict(mapping)
    used = [index for index in mapping.values() if index is not None]
    if (any(type(index) is not int or not 0 <= index < len(headers) for index in used)
            or len(used) != len(set(used))):
        raise UploadError('Invalid or duplicate registration column mapping')
    raw = {field: (values[mapping[field]] if mapping.get(field) is not None else '')
           for field in IMPORT_FIELDS}
    return {
        'row_key': row_key, **raw,
        'normalized_person_name': normalize_text(raw['person_name']),
        'normalized_email': normalize_email(raw['email']),
        'normalized_phone': normalize_phone(raw['phone']),
        'normalized_name': normalize_name(raw['company_name']),
        'normalized_tax_number': normalize_tax(raw['tax_number']),
        'normalized_registration_number': normalize_registration(raw['registration_number']),
        'normalized_address': normalize_text(raw['address']),
        'normalized_municipality': normalize_text(raw['municipality']),
    }


def _validate(headers, rows, max_rows, max_columns=100):
    if not headers or not any(_plain(value) for value in headers):
        raise UploadError('The uploaded file has no header row')
    if len(headers) > max_columns:
        raise UploadError(f'The upload exceeds the {max_columns}-column limit')
    if len(rows) > max_rows:
        raise UploadError(f'The upload exceeds the {max_rows}-row limit')
    if not rows:
        raise UploadError('The uploaded file has no company rows')
    width = len(headers)
    if any(len(row) != width for row in rows):
        raise UploadError('Malformed file: rows have inconsistent column counts')
    return [_plain(value) for value in headers], [[_plain(value) for value in row] for row in rows]


def parse_csv(data, max_rows):
    decoded = None
    for encoding in ('utf-8-sig', 'cp1250'):
        try:
            decoded = data.decode(encoding)
            break
        except UnicodeDecodeError:
            pass
    if decoded is None:
        raise UploadError('CSV must use UTF-8 or Windows-1250 encoding')
    sample = decoded[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
    except csv.Error as exc:
        raise UploadError('Could not determine CSV delimiter') from exc
    try:
        parsed = list(csv.reader(io.StringIO(decoded, newline=''), dialect, strict=True))
    except csv.Error as exc:
        raise UploadError(f'Malformed CSV: {exc}') from exc
    if not parsed:
        raise UploadError('The uploaded file is empty')
    return (*_validate(parsed[0], parsed[1:], max_rows), None)


def validate_xlsx_archive(path):
    """Reject malformed or disproportionately expandable XLSX containers."""
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if not entries:
                raise UploadError('Invalid XLSX workbook')
            if len(entries) > XLSX_MAX_ZIP_ENTRIES:
                raise UploadError('The XLSX file is too complex to process safely')
            total_size = 0
            names = set()
            for entry in entries:
                if entry.filename in names or entry.flag_bits & 0x1:
                    raise UploadError('Invalid XLSX workbook')
                names.add(entry.filename)
                if entry.file_size < 0 or entry.compress_size < 0:
                    raise UploadError('Invalid XLSX workbook')
                if entry.file_size > XLSX_MAX_ENTRY_BYTES:
                    raise UploadError('The XLSX file is too complex to process safely')
                total_size += entry.file_size
                if total_size > XLSX_MAX_UNCOMPRESSED_BYTES:
                    raise UploadError('The XLSX file is too complex to process safely')
                if (entry.file_size >= XLSX_COMPRESSION_RATIO_MIN_BYTES and
                        (entry.compress_size == 0 or
                         entry.file_size / entry.compress_size > XLSX_MAX_COMPRESSION_RATIO)):
                    raise UploadError('The XLSX file is too complex to process safely')
            if archive.testzip() is not None:
                raise UploadError('Invalid XLSX workbook')
    except UploadError:
        raise
    except (OSError, EOFError, RuntimeError, NotImplementedError,
            zipfile.BadZipFile, zipfile.LargeZipFile, zlib.error) as exc:
        raise UploadError('Invalid XLSX workbook') from exc


def parse_xlsx(path, max_rows):
    validate_xlsx_archive(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise UploadError('Invalid XLSX workbook') from exc
    try:
        visible = [sheet for sheet in workbook.worksheets if sheet.sheet_state == 'visible']
        if not visible:
            raise UploadError('The workbook has no visible worksheet')
        sheet = visible[0]
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator, None)
        if headers is None:
            raise UploadError('The worksheet is empty')
        headers = list(headers)
        while headers and not _plain(headers[-1]):
            headers.pop()
        if not headers:
            raise UploadError('The uploaded file has no header row')
        rows = []
        for row in iterator:
            values = list(row)
            if any(value is not None and _plain(value) for value in values[len(headers):]):
                raise UploadError('Malformed file: a row exceeds the header width')
            values = values[:len(headers)]
            values.extend([None] * (len(headers)-len(values)))
            rows.append(values)
            if len(rows) > max_rows:
                raise UploadError(f'The upload exceeds the {max_rows}-row limit')
        # Preserve blank rows inside the registration data so worksheet row
        # numbers remain stable, but ignore unused blank rows after the data.
        while rows and not any(value is not None and _plain(value) for value in rows[-1]):
            rows.pop()
        headers, rows = _validate(headers, rows, max_rows)
        return headers, rows, sheet.title
    finally:
        workbook.close()


def store_and_parse(file_storage, storage_root, max_bytes, max_rows):
    original = Path(file_storage.filename or '').name
    extension = Path(original).suffix.lower()
    if extension not in ('.csv', '.xlsx'):
        raise UploadError('Only .csv and .xlsx files are accepted')
    data = file_storage.stream.read(max_bytes + 1)
    if len(data) > max_bytes:
        raise UploadError(f'The file exceeds the {max_bytes // (1024*1024)} MB limit')
    if not data:
        raise UploadError('The uploaded file is empty')
    upload_id = uuid.uuid4().hex
    root = Path(storage_root).expanduser().absolute()
    directory = root/'uploads'
    directory.mkdir(parents=True, exist_ok=True, mode=0o700)
    if root.is_symlink() or directory.is_symlink():
        raise UploadError('Upload storage may not use symbolic links')
    filename = f'{upload_id}{extension}'
    path = (directory/filename).absolute()
    if directory not in path.parents:
        raise UploadError('Invalid upload storage path')
    descriptor = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL | getattr(os, 'O_NOFOLLOW', 0), 0o600)
    try:
        with os.fdopen(descriptor, 'wb') as target:
            target.write(data)
        headers, rows, sheet = parse_csv(data, max_rows) if extension == '.csv' else parse_xlsx(path, max_rows)
    except BaseException:
        path.unlink(missing_ok=True)
        raise
    return {
        'upload_id': upload_id, 'original_filename': original, 'relative_path': f'uploads/{filename}',
        'sha256': hashlib.sha256(data).hexdigest(), 'size_bytes': len(data),
        'format': extension[1:].upper(), 'worksheet_name': sheet,
        'headers': headers, 'rows': rows,
    }
