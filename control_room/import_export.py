"""Lossless XLSX result export for durable Import & Enrich jobs."""
from collections import Counter
from copy import copy
import hashlib
import os
from pathlib import Path
import tempfile
import uuid
from urllib.parse import urlsplit, urlunsplit

from openpyxl import Workbook, load_workbook

from control_room.import_outputs import decode_requested_outputs


EXPORT_HEADERS = (
    'Dastabase Row ID',
    'Match Status',
    'Match Confidence',
    'Canonical Company ID',
    'Matched Company',
    'Website',
    'Default Email',
    'Phone',
    'Revenue 2025',
    'Profit 2025',
    'Employees 2025',
    'Address',
    'Registration Number',
    'Tax Number',
    'Enrichment / Knowledge Source',
    'Identity Status',
    'Research Scope',
)

MERLIN_METADATA_HEADERS = ('Merlin match', 'Merlin note')
MERLIN_OUTPUT_HEADERS = {
    'WEBSITE': (('Website', 'official_website'),),
    'EMAIL': (('Email', 'default_email'),),
    'PHONE': (('Phone', 'default_phone'),),
    'FINANCIALS': (
        ('Revenue 2025', 'revenue_2025'),
        ('Profit 2025', 'profit_2025'),
        ('Employees 2025', 'employees_2025'),
        ('Assets 2025', 'assets_2025'),
        ('Capital 2025', 'capital_2025'),
    ),
}


def _inside(root, relative_path):
    path = (root / relative_path).resolve()
    try:
        path.relative_to(root.resolve())
    except ValueError as exc:
        raise ValueError('Stored workbook path escapes the Control Room storage root') from exc
    return path


def _sha256(path):
    digest = hashlib.sha256()
    with path.open('rb') as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def _unique_label(base, existing, owner='Dastabase'):
    key = base.strip().casefold()
    if key not in existing:
        existing.add(key)
        return base
    number = 1
    while True:
        suffix = f' ({owner})' if number == 1 else f' ({owner} {number})'
        candidate = f'{base}{suffix}'
        key = candidate.strip().casefold()
        if key not in existing:
            existing.add(key)
            return candidate
        number += 1


def _unique_sheet_title(workbook, base='Dastabase QC'):
    existing = set(workbook.sheetnames)
    if base not in existing:
        return base
    number = 2
    while f'{base} ({number})' in existing:
        number += 1
    return f'{base} ({number})'


def _confidence(item):
    strengths = [entry.get('strength') for entry in item['match_evidence']
                 if isinstance(entry, dict) and entry.get('strength')]
    priority = {'DECISIVE': 3, 'STRONG': 2, 'SUPPORTING': 1}
    return max(strengths, key=lambda value: priority.get(value, 0)) if strengths else ''


def _export_values(job_id, item, research_scope):
    data = item['enrichment'] if item['match_status'] == 'MATCHED' else {}
    sources = data.get('sources') or ()
    return (
        f'{job_id}:{item["item_position"]}',
        item['match_status'],
        _confidence(item),
        item['company_id'] if item['match_status'] == 'MATCHED' else '',
        data.get('canonical_company_name') or '',
        data.get('official_website') or '',
        data.get('default_email') or '',
        data.get('default_phone') or '',
        data.get('revenue_2025') if data.get('revenue_2025') is not None else '',
        data.get('profit_2025') if data.get('profit_2025') is not None else '',
        data.get('employees_2025') if data.get('employees_2025') is not None else '',
        data.get('address') or '',
        data.get('registration_number') or '',
        data.get('tax_number') or '',
        ' | '.join(str(value) for value in sources),
        item.get('identity_status') or '',
        research_scope or '',
    )


def _merlin_headers(requested_outputs):
    return MERLIN_METADATA_HEADERS + tuple(
        header for output in requested_outputs
        for header, _ in MERLIN_OUTPUT_HEADERS[output])


def _merlin_match(item):
    if item['match_status'] == 'MATCHED':
        return 'Matched', ''
    if item['match_status'] == 'AMBIGUOUS':
        return 'Ambiguous', 'Multiple possible company matches. Please review this row.'
    return 'Not found', 'No matching company was found.'


def canonical_merlin_website(value):
    """Project an accepted HTTP(S) evidence URL as its customer-facing origin."""
    if not isinstance(value, str) or not value or value != value.strip():
        return value
    if any(character.isspace() for character in value):
        return value
    try:
        parsed = urlsplit(value)
        if (parsed.scheme not in ('http', 'https') or not parsed.netloc
                or parsed.username is not None or parsed.password is not None
                or not parsed.hostname):
            return value
        parsed.port  # Validate a supplied port before projecting the URL.
    except ValueError:
        return value
    return urlunsplit((parsed.scheme, parsed.netloc, '/', '', ''))


def _merlin_export_values(item, requested_outputs):
    match, note = _merlin_match(item)
    data = item['enrichment'] if item['match_status'] == 'MATCHED' else {}
    values = [match, note]
    for output in requested_outputs:
        for _, key in MERLIN_OUTPUT_HEADERS[output]:
            value = data.get(key)
            if key == 'official_website':
                value = canonical_merlin_website(value)
            values.append(value if value is not None else '')
    return tuple(values)


def _qc_metrics(items, upload_rows):
    statuses = Counter(item['match_status'] for item in items)
    fingerprints = Counter(tuple(str(value) for value in row['original_values'])
                           for row in upload_rows)
    matched = [item for item in items if item['match_status'] == 'MATCHED']
    external = sum(item.get('identity_status') == 'RESOLVED_NEW_ENTITY' for item in items)
    return (
        ('Source registration rows', len(upload_rows)),
        ('Exported registration rows', len(items)),
        ('Matched rows', statuses['MATCHED']),
        ('Ambiguous rows', statuses['AMBIGUOUS']),
        ('Unresolved rows', statuses['UNRESOLVED']),
        ('Skipped / not eligible rows', sum(item.get('identity_status') == 'NOT_ELIGIBLE'
                                            for item in items)),
        ('Duplicate source row groups', sum(count > 1 for count in fingerprints.values())),
        ('Duplicate source rows preserved', sum(max(0, count - 1)
                                                for count in fingerprints.values())),
        ('Rows with website', sum(bool(item['enrichment'].get('official_website'))
                                  for item in matched)),
        ('Rows with default email', sum(bool(item['enrichment'].get('default_email'))
                                        for item in matched)),
        ('Rows with phone', sum(bool(item['enrichment'].get('default_phone'))
                                for item in matched)),
        ('Rows with financials', sum(any(item['enrichment'].get(key) is not None
            for key in ('revenue_2025', 'profit_2025', 'employees_2025')) for item in matched)),
        ('Externally identified new entities', external),
    )


def _merlin_qc_metrics(items, upload_rows):
    statuses = Counter(_merlin_match(item)[0] for item in items)
    fingerprints = Counter(tuple(str(value) for value in row['original_values'])
                           for row in upload_rows)
    return (
        ('Source rows', len(upload_rows)),
        ('Exported rows', len(items)),
        ('Matched rows', statuses['Matched']),
        ('Ambiguous rows', statuses['Ambiguous']),
        ('Not found rows', statuses['Not found']),
        ('Duplicate source row groups', sum(count > 1 for count in fingerprints.values())),
        ('Duplicate source rows preserved', sum(max(0, count - 1)
                                                for count in fingerprints.values())),
    )


def _set_literal(cell, value):
    """Write persisted CSV text without turning formula-looking input into code."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'


def _csv_workbook(upload, upload_rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Merlin Results'
    for column, header in enumerate(upload['headers'], 1):
        _set_literal(sheet.cell(1, column), header)
    for expected, row in enumerate(upload_rows, 2):
        if row['row_number'] != expected or len(row['original_values']) != len(upload['headers']):
            workbook.close()
            raise ValueError('Stored CSV rows are incomplete or malformed')
        for column, value in enumerate(row['original_values'], 1):
            _set_literal(sheet.cell(expected, column), value)
    return workbook, sheet


def export_import_xlsx(repository, job_id, storage_root):
    """Create/replace a lossless legacy or request-aware MERLIN workbook."""
    root = Path(storage_root).expanduser().absolute()
    job = repository.get_job(job_id)
    if not job or job['module'] != 'IMPORT_ENRICH':
        raise ValueError('XLSX export requires an IMPORT_ENRICH job')
    items = repository.job_items(job_id)
    if not items or any(item['processing_status'] != 'COMPLETED' for item in items):
        raise ValueError('All registration rows must be processed before export')
    upload_ids = {item['upload_id'] for item in items}
    if len(upload_ids) != 1:
        raise ValueError('Import export requires exactly one source upload')
    upload = repository.get_upload(next(iter(upload_ids)))
    is_merlin = job.get('origin_surface') == 'MERLIN'
    requested_outputs = (decode_requested_outputs(job.get('requested_outputs_json'))
                         if is_merlin else None)
    if is_merlin and requested_outputs is None:
        raise ValueError('MERLIN export requires requested outputs')
    if not upload or upload['format'] not in ('XLSX', 'CSV'):
        raise ValueError('Import result export requires an XLSX or CSV upload')
    if not is_merlin and upload['format'] != 'XLSX':
        raise ValueError('Legacy Import result export requires an XLSX upload')

    if upload['format'] == 'XLSX':
        if not upload['worksheet_name']:
            raise ValueError('Stored source worksheet is missing')
        source = _inside(root, upload['relative_path'])
        if not source.is_file() or source.is_symlink():
            raise ValueError('Stored source workbook is missing or unsafe')
        if _sha256(source) != upload['sha256']:
            raise ValueError('Stored source workbook hash does not match upload metadata')

    directory = root / 'jobs' / job_id
    directory.mkdir(parents=True, exist_ok=True, mode=0o700)
    if root.is_symlink() or directory.is_symlink():
        raise ValueError('Control Room export storage may not use symbolic links')

    upload_rows = repository.upload_rows(upload['upload_id'])
    if upload['format'] == 'XLSX':
        workbook = load_workbook(source, read_only=False, data_only=False, keep_links=True)
        if upload['worksheet_name'] not in workbook.sheetnames:
            workbook.close()
            raise ValueError('Stored source worksheet is missing')
        sheet = workbook[upload['worksheet_name']]
    else:
        workbook, sheet = _csv_workbook(upload, upload_rows)

    try:
        source_last_column = sheet.max_column
        existing = {str(sheet.cell(1, column).value).strip().casefold()
                    for column in range(1, source_last_column + 1)
                    if sheet.cell(1, column).value is not None}
        base_headers = (_merlin_headers(requested_outputs)
                        if is_merlin else EXPORT_HEADERS)
        owner = 'Merlin' if is_merlin else 'Dastabase'
        labels = [_unique_label(header, existing, owner) for header in base_headers]
        for offset, label in enumerate(labels, source_last_column + 1):
            cell = sheet.cell(1, offset, label)
            if source_last_column:
                template = sheet.cell(1, source_last_column)
                cell._style = copy(template._style)
                cell.number_format = template.number_format
                cell.alignment = copy(template.alignment)
                cell.protection = copy(template.protection)

        expected_rows = {row['row_number'] for row in upload_rows}
        actual_rows = [item['upload_row_number'] for item in items]
        if (len(actual_rows) != len(set(actual_rows)) or set(actual_rows) != expected_rows
                or any(row_number < 2 or row_number > sheet.max_row
                       for row_number in actual_rows)):
            raise ValueError('Import job has invalid, duplicate, or out-of-range source row identity')

        task_scope = ({task['task_id']: task.get('enrichment_scope') or ''
                       for task in repository.identity_tasks(job_id)}
                      if not is_merlin else {})
        for item in items:
            row_number = item['upload_row_number']
            values = (_merlin_export_values(item, requested_outputs) if is_merlin else
                      _export_values(job_id, item,
                                     task_scope.get(item.get('identity_task_id'))))
            for offset, value in enumerate(values, source_last_column + 1):
                sheet.cell(row_number, offset, value)

        qc_title = 'Merlin review' if is_merlin else 'Dastabase QC'
        qc = workbook.create_sheet(_unique_sheet_title(workbook, qc_title))
        qc.append((('Merlin Import & Enrich review' if is_merlin else
                    'Dastabase Import & Enrich QC'), 'Value'))
        if not is_merlin:
            qc.append(('Job ID', job_id))
        metrics = (_merlin_qc_metrics(items, upload_rows) if is_merlin else
                   _qc_metrics(items, upload_rows))
        for metric in metrics:
            qc.append(metric)
        qc.append(())
        qc.append(('Export column', 'Workbook header'))
        for base, actual in zip(base_headers, labels):
            qc.append((base, actual))
        qc.freeze_panes = 'A2'
        qc.column_dimensions['A'].width = 38
        qc.column_dimensions['B'].width = 32

        descriptor, temporary_name = tempfile.mkstemp(
            prefix='.import-enriched-', suffix='.xlsx', dir=directory)
        os.close(descriptor)
        temporary = Path(temporary_name)
        try:
            workbook.save(temporary)
            os.chmod(temporary, 0o600)
        except BaseException:
            temporary.unlink(missing_ok=True)
            raise
    finally:
        workbook.close()

    try:
        size = temporary.stat().st_size
        digest = _sha256(temporary)
        output = _inside(root, Path('jobs') / job_id /
                         f'import-enriched-{uuid.uuid4().hex}.xlsx')
        # link() publishes a fully written same-filesystem inode and refuses to
        # replace an existing immutable generation.
        os.link(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)
    relative = output.relative_to(root).as_posix()
    repository.add_artifact(job_id, 'UPLOAD_RECONCILIATION', relative, size, digest)
    return {'path': output, 'relative_path': relative,
            'headers': dict(zip(base_headers, labels)), 'qc': dict(metrics)}
